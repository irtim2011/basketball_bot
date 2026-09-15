"""End-to-end boundaries for forecasts, primary answers and workbook delivery.

All Telegram traffic and the workbook are local fakes; no live users or Sheets
are contacted. The database, scheduler and answer handlers are the real modules.
"""
import asyncio
from datetime import datetime, timedelta
from pathlib import Path
from tempfile import TemporaryDirectory
from types import SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

import background
import db
import events
import google_sheet
import handlers_poll
import scheduler
import utils


class FakeMessage:
    def __init__(self, bot, message_id, chat_id, text, reply_markup=None):
        self.bot = bot
        self.message_id = message_id
        self.chat = SimpleNamespace(id=chat_id, type='private')
        self.from_user = SimpleNamespace(id=chat_id)
        self.text = text
        self.reply_markup = reply_markup
        self.edits = []

    async def edit_text(self, text, reply_markup=None, **kwargs):
        self.text, self.reply_markup = text, reply_markup
        self.edits.append(text)
        return self

    async def answer(self, text, **kwargs):
        return await self.bot.send_message(self.chat.id, text, **kwargs)


class FakeTelegram:
    id = 999999

    def __init__(self):
        self.messages = []

    async def send_message(self, chat_id, text, reply_markup=None, **kwargs):
        message = FakeMessage(self, len(self.messages) + 1, chat_id, text, reply_markup)
        self.messages.append(message)
        return message

    def primary_messages(self):
        return [message for message in self.messages
                if message.reply_markup and any(
                    (button.callback_data or '').startswith('r:')
                    for row in message.reply_markup.inline_keyboard for button in row)]


class PlanningIntegrationTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.tmp = TemporaryDirectory()
        self.old_path = db.DB_PATH
        db.DB_PATH = str(Path(self.tmp.name) / 'planning-integration.sqlite3')
        self.now = utils.TZ.localize(datetime(2026, 9, 20, 11, 59))  # Before weekly launch.
        self.clock_patch = patch('utils.now', return_value=self.now)
        self.clock = self.clock_patch.start()
        self.queue_patch = patch('google_sheet.queue', return_value=True)
        self.queue = self.queue_patch.start()
        self.delivery_patch = patch.object(scheduler, 'delivery_lock', asyncio.Lock())
        self.delivery_patch.start()
        self.offset_patch = patch.object(scheduler, 'POLL_OFFSET_MINUTES', 1440)
        self.offset_patch.start()
        self.addAsyncCleanup(self.cleanup_resources)
        await db.init_db()
        self.pid = await db.register_participant(77, 'dima', 'Васильев Дмитрий', '+79991234567')
        await db.set_active(self.pid, True)
        self.bot = FakeTelegram()

    async def cleanup_resources(self):
        await background.close()
        await google_sheet.close()
        await db.close_db()
        db._conn = None
        self.offset_patch.stop()
        self.delivery_patch.stop()
        self.queue_patch.stop()
        self.clock_patch.stop()
        db.DB_PATH = self.old_path
        self.tmp.cleanup()

    async def create_slot(self, iso='2026-09-22T19:00:00'):
        start = utils.TZ.localize(datetime.fromisoformat(iso))
        sid = await events.save_slot(start.weekday(), start.strftime('%H:%M'),
                                     start.date().isoformat(), starts_on=start.date().isoformat())
        return sid, start

    async def response(self, sid):
        return await (await db._c().execute(
            'SELECT * FROM responses WHERE participant_id=? AND schedule_id=?',
            (self.pid, sid))).fetchone()

    async def callback(self, message, response_id, answer, token, at):
        self.clock.return_value = at
        callback = SimpleNamespace(id=token, from_user=SimpleNamespace(id=77),
                                   data=f'r:{response_id}:{answer}', message=message,
                                   answer=AsyncMock())
        await handlers_poll.process_answer(callback)
        callback.answer.assert_awaited_once()
        return callback

    async def test_primary_opens_at_exact_24h_and_does_not_duplicate_after_restart(self):
        self.assertEqual(scheduler.POLL_OFFSET_MINUTES, 1440)
        sid, start = await self.create_slot()
        self.clock.return_value = start - timedelta(hours=24, seconds=1)
        await scheduler.tick(self.bot)
        self.assertEqual(self.bot.primary_messages(), [])
        self.assertIsNone(await self.response(sid))

        self.clock.return_value = start - timedelta(hours=24)
        await scheduler.tick(self.bot)
        self.assertEqual(len(self.bot.primary_messages()), 1)
        before = await self.response(sid)
        await scheduler.tick(self.bot)
        await db.close_db()
        await db.init_db()
        await scheduler.tick(self.bot)
        after = await self.response(sid)
        self.assertEqual(len(self.bot.primary_messages()), 1)
        self.assertEqual((after['id'], after['message_id']), (before['id'], before['message_id']))

    async def test_same_primary_buttons_change_financial_mark_and_work_after_start(self):
        import planning
        sid, start = await self.create_slot()
        self.clock.return_value = start - timedelta(hours=24)
        await scheduler.tick(self.bot)
        response = await self.response(sid)
        message = self.bot.primary_messages()[0]
        initial_buttons = [[button.callback_data for button in row]
                           for row in message.reply_markup.inline_keyboard]
        self.queue.reset_mock()
        transitions = [('no', 'N', start - timedelta(hours=23)),
                       ('yes', 'Y', start - timedelta(hours=22)),
                       ('no', 'N', start - timedelta(hours=21)),
                       ('yes', 'Y', start + timedelta(minutes=15))]
        for index, (answer, mark, at) in enumerate(transitions):
            await self.callback(message, response['id'], answer, f'main-change-{index}', at)
            self.assertEqual((await self.response(sid))['status'], answer)
            _, rows = await events.summary()
            participant = next(row for row in rows if row['telegram_id'] == 77)
            self.assertEqual(participant['marks'][start.date().isoformat()], mark)
            self.assertEqual([[button.callback_data for button in row]
                              for row in message.reply_markup.inline_keyboard], initial_buttons)
        self.assertEqual(self.queue.call_count, len(transitions))
        stage_log = [tuple(row) for row in await (await db._c().execute(
            "SELECT old_status,new_status FROM answer_change_log "
            "WHERE log_type='stage' AND to_stage='main' ORDER BY id")).fetchall()]
        self.assertEqual(stage_log, [('pending', 'no'), ('no', 'yes'), ('yes', 'no'), ('no', 'yes')])
        # A delayed replay of the first callback must neither undo the answer
        # nor display that obsolete N while the stored answer is Y.
        await self.callback(message, response['id'], 'no', 'main-change-0', start + timedelta(minutes=16))
        self.assertEqual((await self.response(sid))['status'], 'yes')
        import texts
        self.assertEqual(message.text, texts.poll_text(start, 'yes'))
        replay_log = [tuple(row) for row in await (await db._c().execute(
            "SELECT old_status,new_status FROM answer_change_log "
            "WHERE log_type='stage' AND to_stage='main' ORDER BY id")).fetchall()]
        self.assertEqual(replay_log, stage_log)
        # The public planning snapshot and the actual answer survive restarting.
        before = await planning.snapshot(self.clock.return_value)
        await db.close_db()
        await db.init_db()
        after = await planning.snapshot(self.clock.return_value)
        self.assertEqual(before, after)
        self.assertEqual((await self.response(sid))['status'], 'yes')
        persisted = [tuple(row) for row in await (await db._c().execute(
            "SELECT old_status,new_status FROM answer_change_log "
            "WHERE log_type='stage' AND to_stage='main' ORDER BY id")).fetchall()]
        self.assertEqual(persisted, stage_log)

    async def test_weekly_forecast_yes_never_becomes_financial_attendance(self):
        import planning
        import handlers_planning
        await self.create_slot('2026-09-24T19:00:00')
        await planning.deliver_due(self.bot, self.clock.return_value)
        self.assertEqual(self.bot.messages, [])
        self.clock.return_value = utils.TZ.localize(datetime(2026, 9, 20, 12, 0))
        await planning.deliver_due(self.bot, self.clock.return_value)
        row = await (await db._c().execute(
            'SELECT a.*, p.message_id FROM planned_answers a '
            'JOIN planned_polls p ON p.id=a.poll_id ORDER BY a.id LIMIT 1')).fetchone()
        self.assertIsNotNone(row)
        message = next(message for message in self.bot.messages if message.message_id == row['message_id'])
        callback = SimpleNamespace(id='weekly-forecast-yes', from_user=SimpleNamespace(id=77),
            message=message, answer=AsyncMock(), data=f"pa:{row['id']}:{row['generation']}:yes:0")
        self.queue.reset_mock()
        await handlers_planning.answer(callback)
        callback.answer.assert_awaited_once()
        self.queue.assert_called_once()
        saved = await (await db._c().execute('SELECT status FROM planned_answers WHERE id=?',
                                             (row['id'],))).fetchone()
        self.assertEqual(saved['status'], 'yes')
        for table in ('responses', 'attendance'):
            count = await (await db._c().execute(f'SELECT COUNT(*) FROM {table}')).fetchone()
            self.assertEqual(count[0], 0)
        _, rows = await events.summary()
        self.assertFalse(any(mark for row in rows for mark in row['marks'].values()))
        await planning.deliver_due(self.bot, self.clock.return_value)
        count = await (await db._c().execute('SELECT COUNT(*) FROM planned_polls')).fetchone()
        self.assertEqual(count[0], 1)

    async def test_monthly_forecast_starts_only_at_noon_on_last_day_and_is_not_resent(self):
        import planning
        await self.create_slot('2026-10-05T19:00:00')
        self.clock.return_value = utils.TZ.localize(datetime(2026, 9, 30, 11, 59, 59))
        await planning.deliver_due(self.bot, self.clock.return_value)
        count = await (await db._c().execute("SELECT COUNT(*) FROM planned_polls WHERE kind='month'")).fetchone()
        self.assertEqual(count[0], 0)
        self.clock.return_value += timedelta(seconds=1)
        await planning.deliver_due(self.bot, self.clock.return_value)
        row = await (await db._c().execute("SELECT * FROM planned_polls WHERE kind='month'")).fetchone()
        self.assertIsNotNone(row)
        self.assertEqual(row['period_start'], '2026-10-01')
        self.assertEqual(row['period_end'], '2026-10-31')
        self.assertEqual(row['status'], 'sent')
        message_count = len(self.bot.messages)
        await db.close_db()
        await db.init_db()
        await planning.deliver_due(self.bot, self.clock.return_value)
        self.assertEqual(len(self.bot.messages), message_count)

    async def test_primary_updates_do_not_overwrite_manual_attendance_correction(self):
        import attendance_sync
        from test_attendance_sync import FakeBook
        book = FakeBook()
        participant = await db.get_participant(self.pid)
        for sheet in book.sheets.values():
            sheet.values[2][0] = str(participant['public_id'])
        sid, start = await self.create_slot()
        self.clock.return_value = start - timedelta(hours=24)
        await scheduler.tick(self.bot)
        response = await self.response(sid)
        message = self.bot.primary_messages()[0]
        dates, rows = await events.summary()
        attendance_sync.reconcile_book(book, dates, rows)  # Baseline before first actual answer.

        await self.callback(message, response['id'], 'yes', 'first-real-yes', start - timedelta(hours=23))
        dates, rows = await events.summary()
        attendance_sync.reconcile_book(book, dates, rows)
        column = book.sheets['Посещения'].values[0].index(start.strftime('%d.%m.%Y'))
        self.assertEqual(book.sheets['Посещения'].values[2][column], 'Y')
        book.sheets['Посещения'].values[2][column] = ''
        attendance_sync.reconcile_book(book, dates, rows)
        for index, answer in enumerate(('no', 'yes')):
            await self.callback(message, response['id'], answer, f'after-manual-{index}',
                                start + timedelta(minutes=index + 1))
            dates, rows = await events.summary()
            attendance_sync.reconcile_book(book, dates, rows)
            self.assertEqual(book.sheets['Посещения'].values[2][column], '')
            self.assertEqual(book.sheets['Посещения_bot'].values[2][column], '')

    async def test_table_uses_complete_google_workbook_export_with_new_planning_tabs(self):
        import handlers_trainer
        from openpyxl import Workbook, load_workbook
        exported = Path(self.tmp.name) / 'complete-google-book.xlsx'
        workbook = Workbook()
        workbook.active.title = 'Посещения'
        expected = {'Посещения', 'Тарифы', 'План тренировок', 'Состав тренировки', 'История ответов'}
        for title in sorted(expected - {'Посещения'}):
            workbook.create_sheet(title)
        workbook.save(exported)
        workbook.close()
        received = []

        async def capture(document, **kwargs):
            saved = load_workbook(document.path, read_only=True)
            try:
                received.extend(saved.sheetnames)
            finally:
                saved.close()

        message = SimpleNamespace(from_user=SimpleNamespace(id=77),
                                  chat=SimpleNamespace(id=77, type='private'),
                                  answer=AsyncMock(), answer_document=AsyncMock(side_effect=capture))
        with patch.object(handlers_trainer, 'TRAINER_IDS', {77}), \
             patch('google_sheet.configured', return_value=True), \
             patch('google_sheet.sync_now', new_callable=AsyncMock), \
             patch('google_sheet.export_workbook_xlsx', return_value=str(exported)) as cloud_export, \
             patch('handlers_trainer.build_xlsx') as attendance_only_export:
            await handlers_trainer.export_table(message)
        cloud_export.assert_called_once()
        attendance_only_export.assert_not_called()
        self.assertEqual(set(received), expected)
        self.assertFalse(exported.exists())


if __name__ == '__main__':
    unittest.main()
